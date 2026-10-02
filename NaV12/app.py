"""NaV12: single-instance Flask application with locked Excel storage.

Linux/macOS only (fcntl); deploy one worker on persistent storage.
"""
import os
import re
import secrets
import fcntl
import threading
import time
from collections import deque
from contextlib import contextmanager
from zipfile import BadZipFile
from datetime import timedelta
from pathlib import Path
from functools import wraps
from flask import Flask, abort, flash, redirect, render_template, request, session, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__)
production = os.environ.get('APP_ENV', 'development') == 'production'
secret = os.environ.get('SECRET_KEY')
if production and (not secret or len(secret) < 32):
    raise RuntimeError('Set SECRET_KEY to a random value of at least 32 characters.')
# Restrict newly created workbook/lock files to the service owner.
os.umask(0o077)
app.config.update(
    SECRET_KEY=secret or secrets.token_hex(32),
    USER_FILE=Path(os.environ.get('DATA_DIR', str(BASE_DIR / 'data'))) / 'users.xlsx',
    SESSION_COOKIE_SECURE=production,
    MAX_USERS=int(os.environ.get('MAX_USERS', '1000')),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    PERMANENT_SESSION_LIFETIME=timedelta(hours=2),
    MAX_CONTENT_LENGTH=16 * 1024,
)


storage_mutex = threading.Lock()
rate_mutex = threading.Lock()
rate_buckets = {}


class StorageUnavailable(Exception):
    pass


@contextmanager
def workbook_lock(path):
    # Thread lock + OS lock: protect a complete read/check/append/save transaction.
    with storage_mutex:
        with open(str(path) + '.lock', 'a') as handle:
            deadline = time.monotonic() + 5
            while True:
                try:
                    fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
                    break
                except BlockingIOError:
                    if time.monotonic() >= deadline:
                        raise StorageUnavailable('Workbook is busy')
                    time.sleep(0.05)
            try:
                yield
            finally:
                fcntl.flock(handle, fcntl.LOCK_UN)


def rate_allowed(username):
    # Bounded, in-memory protection for ONE worker. Resets at process restart.
    # Global cap covers username rotation; per-name cap slows repeated attempts.
    now = time.monotonic()
    keys = [('all', 60), ('user:' + username.casefold(), 10)]
    with rate_mutex:
        for key in list(rate_buckets):
            bucket = rate_buckets[key]
            while bucket and bucket[0] <= now - 60:
                bucket.popleft()
            if not bucket:
                del rate_buckets[key]
        if any(len(rate_buckets.get(key, ())) >= limit for key, limit in keys):
            return False
        for key, _ in keys:
            rate_buckets.setdefault(key, deque()).append(now)
        return True


def workbook_path():
    path = Path(app.config['USER_FILE'])
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_workbook(wb, path):
    # Replace only after a complete write, so an interrupted save keeps the old file.
    temporary = path.with_suffix('.tmp.xlsx')
    try:
        wb.save(temporary)
        with temporary.open('rb') as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def open_users(path):
    if path.exists():
        wb = load_workbook(path)
        if 'Users' not in wb.sheetnames or [c.value for c in wb['Users'][1]] != ['Username', 'Password']:
            wb.close()
            raise StorageUnavailable('Unexpected workbook format; restore a valid backup')
        return wb
    wb = Workbook()
    ws = wb.active
    ws.title = 'Users'
    ws.append(['Username', 'Password'])
    ws.freeze_panes = 'A2'
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 90
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='172126')
    save_workbook(wb, path)
    return wb


def account(username, password=None):
    """Read one hash, or append an account; lock the entire read/write operation."""
    path = workbook_path()
    with workbook_lock(path):
        wb = open_users(path)
        try:
            ws = wb['Users']
            for name, saved_hash in ws.iter_rows(min_row=2, max_col=2, values_only=True):
                if name and str(name).casefold() == username.casefold():
                    return False if password is not None else saved_hash
            if password is None:
                return None
            if ws.max_row - 1 >= app.config['MAX_USERS']:
                raise StorageUnavailable('Registration capacity reached')
            ws.append([username, generate_password_hash(password)])
            save_workbook(wb, path)
            return True
        finally:
            wb.close()


def csrf_token():
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_hex(32)
    return session['csrf_token']


app.jinja_env.globals['csrf_token'] = csrf_token


@app.before_request
def protect_forms():
    if request.method == 'POST':
        supplied = request.form.get('csrf_token', '')
        expected = session.get('csrf_token', '')
        if not expected or not secrets.compare_digest(supplied, expected):
            abort(400, description='Form expired. Reload the page and try again.')
        if request.endpoint in {'login', 'signup'}:
            username = request.form.get('username', '').strip()
            if len(username) > 30:
                abort(400, description='Username is too long.')
            if not rate_allowed(username):
                abort(429, description='Too many attempts. Please wait one minute.')


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if 'username' not in session:
            flash('Please log in to access your garage.', 'error')
            return redirect(url_for('login'))
        return view(*args, **kwargs)
    return wrapped


@app.route('/')
def index():
    return redirect(url_for('dashboard' if 'username' in session else 'login'))


@app.route('/signup', methods=['GET', 'POST'])
def signup():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        if not re.fullmatch(r'[A-Za-z0-9_]{3,30}', username):
            flash('Use 3–30 letters, numbers, or underscores for your username.', 'error')
        elif not 8 <= len(password) <= 128:
            flash('Your password must contain 8–128 characters.', 'error')
        elif not account(username, password):
            flash('Username already exists. Please log in or choose another.', 'error')
        else:
            flash('Account created. Log in to enter your garage.', 'success')
            return redirect(url_for('login'))
    return render_template('signup.html', title='Join the garage')


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        saved_hash = account(username)
        if saved_hash is None:
            flash('User not found', 'error')
        elif not password or len(password) > 128 or not check_password_hash(saved_hash, password):
            flash('Incorrect credentials', 'error')
        else:
            session.clear()
            session.permanent = True
            session['username'] = username
            flash('Successfully logged in', 'success')
            return redirect(url_for('dashboard'))
    return render_template('login.html', title='Welcome back')


@app.route('/dashboard')
@login_required
def dashboard():
    return render_template('dashboard.html', title='Your garage')


@app.post('/logout')
def logout():
    session.clear()
    flash('You have been logged out.', 'success')
    return redirect(url_for('login'))


@app.after_request
def prevent_cached_accounts(response):
    if request.endpoint != 'static':
        response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Content-Security-Policy'] = (
        "default-src 'self'; style-src 'self' 'unsafe-inline'; "
        "img-src 'self'; script-src 'none'; object-src 'none'; "
        "base-uri 'self'; frame-ancestors 'none'; form-action 'self'"
    )
    if production:
        response.headers['Strict-Transport-Security'] = 'max-age=31536000'
    return response


@app.get('/healthz')
def health():
    return {'status': 'ok'}


@app.errorhandler(StorageUnavailable)
@app.errorhandler(OSError)
@app.errorhandler(BadZipFile)
def storage_error(error):
    app.logger.error('Account storage unavailable: %s', type(error).__name__)
    return render_template('error.html', title='Temporarily unavailable',
                           message='Account storage is temporarily unavailable. Please try again later.'), 503


@app.errorhandler(400)
@app.errorhandler(413)
@app.errorhandler(429)
def request_error(error):
    response = app.make_response((render_template('error.html', title='Request could not be completed',
                                                 message=error.description), error.code))
    if error.code == 429:
        response.headers['Retry-After'] = '60'
    return response


# Runs at import, including under Gunicorn and PythonAnywhere WSGI.
# Never silently overwrite an existing damaged workbook.
with workbook_lock(workbook_path()):
    initial_workbook = open_users(workbook_path())
    initial_workbook.close()


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=False)
