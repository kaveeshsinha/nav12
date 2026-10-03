from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

USERNAME = "laferrari"
NEW_PASSWORD = "1234"

wb = load_workbook("data/users.xlsx")
ws = wb["Users"]

found = False

for row in ws.iter_rows(min_row=2):
    username = row[0].value

    if username and str(username).casefold() == USERNAME.casefold():
        row[1].value = generate_password_hash(NEW_PASSWORD)
        found = True
        break

if found:
    wb.save("data/users.xlsx")
    print("Password reset successfully.")
else:
    print("Username not found.")

wb.close()