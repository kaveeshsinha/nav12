"""Import a private Excel backup into the configured database; dry-run by default."""
import argparse
import os
import re
import sys
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

from accounts import AccountStore, ImportConflict, StorageUnavailable, valid_username


def supported_hash(value):
    if not isinstance(value, str) or len(value) > 512:
        return False
    if not re.fullmatch(r'(scrypt:[0-9]+:[0-9]+:[0-9]+|pbkdf2:sha(256|512):[0-9]+)\$[A-Za-z0-9]+\$[0-9a-f]+', value):
        return False
    method, salt, digest = value.split('$')
    parts = method.split(':')
    if parts[0] == 'scrypt':
        n, r, p = map(int, parts[1:])
        return n > 1 and n & (n - 1) == 0 and r > 0 and p > 0 and len(digest) == 128
    return int(parts[2]) > 0 and len(digest) == (64 if parts[1] == 'sha256' else 128)


def read_backup(path):
    with Path(path).open('rb') as handle:
        workbook = load_workbook(handle, read_only=True, data_only=False)
        try:
            if 'Users' not in workbook.sheetnames:
                raise ValueError('Expected a Users worksheet.')
            sheet = workbook['Users']
            if [cell.value for cell in sheet[1]] != ['Username', 'Password']:
                raise ValueError('Expected Username and Password headers.')
            rows = []
            seen = set()
            for number, (username, password_hash) in enumerate(
                    sheet.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
                if username is None and password_hash is None:
                    continue
                if not valid_username(username):
                    raise ValueError(f'Invalid username in row {number}.')
                if not supported_hash(password_hash):
                    raise ValueError(f'Row {number} does not contain a supported password hash. Plaintext passwords are not imported.')
                key = username.casefold()
                if key in seen:
                    raise ValueError(f'Case-insensitive duplicate username in row {number}.')
                seen.add(key)
                rows.append((username, password_hash))
            return rows
        finally:
            workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', help='Path to a private users.xlsx backup')
    parser.add_argument('--apply', action='store_true', help='Write accounts after validation; otherwise preview only')
    args = parser.parse_args()
    database_url = os.environ.get('DATABASE_URL', '').strip()
    if not database_url:
        parser.error('Set DATABASE_URL to the destination database first.')
    store = None
    try:
        rows = read_backup(args.workbook)
        store = AccountStore(database_url, production=os.environ.get('APP_ENV') == 'production',
                             max_users=int(os.environ.get('MAX_USERS', '1000')))
        store.initialize()
        result = store.import_hashes(rows, apply=args.apply)
    except (ValueError, OSError, BadZipFile, StorageUnavailable, ImportConflict) as error:
        # Known errors contain no usernames, hashes, database URLs, or credentials.
        print(f'Import failed: {error}', file=sys.stderr)
        return 1
    finally:
        if store is not None:
            store.close()
    label = 'Imported' if args.apply else 'Dry run: would import'
    print(f"{label} {result['added']} account(s); {result['skipped']} identical account(s) skipped.")
    if not args.apply:
        print('No accounts were changed. Add --apply to import this backup.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
